#!/usr/bin/env python3
"""
Build a consolidated, de-duplicated inventory of calculators from the offline mirrors
and save it as calculator_inventory.xlsx.

Reads each site's own index/category pages to get calculator names, URLs and the
site's own category, merges the same calculator across sites, and assigns every
finance calculator to one of the subcategories in FINANCE_SUBCATEGORIES.

Re-run any time a mirror is refreshed:
    .venv/bin/python build_calculator_inventory.py
"""
import re
from collections import defaultdict
from pathlib import Path

from bs4 import BeautifulSoup
from openpyxl import Workbook
from datetime import date

from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
OUT_FILE = HERE / "calculator_inventory.xlsx"
SITES = ["calculator.net", "dinkytown.net", "fncalculator.com"]

# --- Finance subcategories -------------------------------------------------
# Ordered: the first rule whose pattern matches the calculator name wins, so
# specific topics come before broad ones (e.g. "401k" before "retirement").
FINANCE_SUBCATEGORIES = [
    ("Retirement Accounts (401k, 403b, 457, IRA, Roth)",
     r"401 ?\(?k|403 ?\(?b|457|72 ?\(?t\b|\bira\b|roth|\brmd\b|required minimum|rrsp|rrif|tfsa|\blira\b|\bsep\b|simple ira|keogh|self-employed retirement|pension|stretch|net unrealized appreciation|\bnua\b"),
    ("Social Security & Annuities",
     r"social security|annuit|\bcpp\b|\boas\b|canada pension|old age security"),
    ("Health Savings & Medical Costs",
     r"\bhsa\b|health savings|\bfsa\b|flexible spending|medical|healthcare|health care|long-term care|long term care"),
    ("Insurance",
     r"insurance|disability|life cover|premium|critical illness|human life value"),
    ("Student Loans & Education Savings",
     r"student|college|education|529|\bresp\b|tuition|school"),
    ("Auto Loans & Leasing",
     r"\bauto\b|\bcar\b|vehicle|lease|motorcycle|\brv\b|boat"),
    ("Home Equity & Refinancing",
     r"home equity|heloc|refinanc|reverse mortgage|cash[- ]out|debt consolidation"),
    ("Real Estate & Rental Property",
     r"rental|rent vs|rent or buy|rent versus|buy vs rent|landlord|real estate|property|cap rate|house flip|down payment assistance|\brent\b|net proceeds"),
    ("Mortgage & Home Buying",
     r"mortgage|amorti[sz]ation|house|home|\bfha\b|\bva\b|\barm\b|adjustable|discount points|points calculator|mortgage points|bi-?weekly|interest only|closing cost|down payment|escrow|piti|balloon"),
    ("Credit Cards & Debt Payoff",
     r"credit card|debt|payoff|pay off|minimum payment|credit|snowball|avalanche|consolidat|\bowe\b"),
    ("Payroll, Salary & Take-home Pay",
     r"payroll|paycheck|pay ?check|salary|hourly|wage|take[- ]home|overtime|bonus|commission|w-4|w4|withholding|raise|net pay|gross pay|employee"),
    ("Income Tax & Tax Planning",
     r"\btax\b|taxes|1040|capital gain|estate|gift|\bamt\b|deduction|marginal|tax-?equivalent|\bvat\b|\bgst\b|\bhst\b"),
    ("Bonds, CDs & Fixed Income",
     r"\bbond|treasury|t-bill|\bcd\b|certificate of deposit|fixed income|yield to|coupon|money market|\bgic\b"),
    ("Stocks, Dividends & Trading",
     r"stock|dividend|option|black-?scholes|pivot|fibonacci|capm|beta|share|futures|future contract|commodit|margin call|\bdrip\b|trading|position size|expected return|holding period|p/e|price.?earnings"),
    ("Business Finance & Valuation",
     r"business|break[- ]?even|markup|margin|wacc|cost of capital|ratio|balance sheet|income statement|cash flow|depreciat|inventory|payback|valuation|profit|revenue|sales forecast|forecast|working capital|lease vs buy equipment|equipment|startup|operating|accounts|payback|\bebit"),
    ("Investment Returns & Growth",
     r"invest|return|\broi\b|\birr\b|\bnpv\b|net present|future value|present value|\bfv\b|\bpv\b|growth|asset allocation|portfolio|mutual fund|fund fee|expense ratio|\betf\b|lump sum|dollar cost|cagr|average return|risk"),
    ("Retirement Planning",
     r"retire|nest egg|withdrawal|distribution|fire\b|early retirement|life expectancy"),
    ("Interest, APR & Time Value of Money",
     r"interest|\bapr\b|\bapy\b|\btvm\b|time value|rule of 72|effective rate|compound|simple interest|discount rate|rate converter|nominal"),
    ("Savings & Emergency Funds",
     r"sav(e|ing)|emergency|goal|millionaire|million|deposit|nest|piggy|christmas club"),
    ("Personal & General Loans",
     r"loan|borrow|payment calculator|payment frequenc|lending|financing|installment"),
    ("Budgeting, Net Worth & Cash Flow",
     r"budget|net worth|cash flow|expense|spending|cost of living|lifestyle|debt-to-income|dti|financial health|allowance|wedding|baby|household|subscription|latte|habit|check ?book|cheque ?book|balancer|spouse"),
    ("Inflation, Currency & Economics",
     r"inflation|currency|exchange|purchasing power|\bcpi\b|gdp|economic"),
    ("Everyday Money & Utility Tools",
     r"\btip\b|discount|sales tax|fuel|gas|percent|coupon|price|shopping|unit price|electric|energy|commute|"
     r"basic|finance calculator|financial calculator|\bdate\b"),
]
FINANCE_SUBCATEGORIES = [(name, re.compile(rx, re.I)) for name, rx in FINANCE_SUBCATEGORIES]

# Checked before FINANCE_SUBCATEGORIES for names whose generic keywords mislead.
OVERRIDES = [(re.compile(rx, re.I), sub) for rx, sub in [
    (r"equipment|business|commercial|debt service coverage|dscr", "Business Finance & Valuation"),
    (r"1031|like kind", "Real Estate & Rental Property"),
    (r"sales tax|\bvat\b|discount and tax", "Everyday Money & Utility Tools"),
    (r"tax[- ]equivalent|certificate", "Bonds, CDs & Fixed Income"),
    (r"^stock", "Stocks, Dividends & Trading"),
    (r"earned income credit|\beic\b", "Income Tax & Tax Planning"),
    (r"civilian pay|military", "Payroll, Salary & Take-home Pay"),
    (r"dealer financing|truck", "Auto Loans & Leasing"),
    (r"equity line", "Home Equity & Refinancing"),
    (r"taxes and insurance|\bpiti\b|\bpmi\b", "Mortgage & Home Buying"),
    (r"thrift savings", "Retirement Accounts (401k, 403b, 457, IRA, Roth)"),
    (r"retirement planner", "Retirement Planning"),
]]
UNCATEGORIZED = "Uncategorized (review)"

# Name variants that mean the same calculator. Keys and values are normalised keys.
ALIASES = {
    "return on investment roi": "roi",
    "return on investment": "roi",
    "home affordablility": "house affordability",
    "home affordability": "house affordability",
    "rent vs buy": "rent vs buy",
    "rent or buy": "rent vs buy",
    "buy vs rent": "rent vs buy",
    "irr npv": "irr npv",
    "compound interest": "compound interest",
    "certificate of deposit cd": "cd",
    "certificate of deposit": "cd",
    "401k contribution": "401k",
    "us inflation": "inflation",
    "us paycheck tax": "paycheck",
    "us health savings account": "hsa",
    "health savings account hsa": "hsa",
    "health savings account": "hsa",
    "black scholes option": "black scholes option",
    "credit card payoff": "credit card payoff",
    "credit card minimum": "credit card minimum payment",
    "credit card minimum payment": "credit card minimum payment",
    "loan refinance": "refinance",
    "mortgage refinance": "refinance",
    "discount points": "mortgage points",
    "mortgage points": "mortgage points",
    "biweekly payment": "biweekly mortgage",
    "bi weekly payment": "biweekly mortgage",
    "biweekly mortgage": "biweekly mortgage",
    "traditional ira vs roth ira": "roth vs traditional ira",
    "roth ira vs traditional ira": "roth vs traditional ira",
    "required minimum distribution rmd": "rmd",
    "required minimum distribution": "rmd",
    "margin and markup": "margin",
    "margin markup": "margin",
    "hourly to salary": "salary hourly",
    "salary to hourly": "salary hourly",
    "hourly salary": "salary hourly",
    "salary": "salary hourly",
    "tax equivalent yield": "tax equivalent yield",
    "rule of 72": "rule of 72",
    "time value of money tvm": "tvm",
    "tvm": "tvm",
    "us treasury bill": "treasury bill",
    "t bill": "treasury bill",
    "fixed vs adjustable rate": "fixed vs adjustable rate mortgage",
    "fixed vs adjustable rate mortgage": "fixed vs adjustable rate mortgage",
    "adjustable rate": "adjustable rate mortgage",
    "adjustable rate mortgage arm": "adjustable rate mortgage",
    "interest only": "interest only mortgage",
    "interest only mortgage": "interest only mortgage",
    "internal rate of return irr": "irr",
    "internal rate of return": "irr",
    "mutual fund expense": "mutual fund",
    "mutual fund fee": "mutual fund",
    "home rent vs buy": "rent vs buy",
    "lease": "auto lease",
    "car loan": "auto loan",
    # Same tool under different names across sites (reviewed by hand).
    "mortgage tax savings": "mortgage tax saving",
    "social security benefit": "social security",
    "municipal bond tax equivalent yield": "tax equivalent yield",
    "arm vs fixed rate mortgage": "fixed vs adjustable rate mortgage",
    "currency converter": "currency",
    "profit margin": "margin",
    "irr npv": "irr",
    "estate tax planning": "estate tax",
    "accelerated debt payoff": "debt payoff",
    "discount and tax": "discount",
    "home budget": "budget",
    "retirement pension": "pension",
    "investment returns": "investment",
    "inflation and consumer prices": "inflation",
    "traditional ira": "ira",
    "health savings account hsa savings": "hsa",
    "credit card": "credit card payoff",
    "amortizing loan": "amortization",
    "take home paycheck": "paycheck",
    "paycheck tax": "paycheck",
    "payroll deductions": "paycheck",
    "1040 tax": "income tax",
    "simple federal tax": "income tax",
    "u s easy tax": "income tax",
    "should i refinance": "refinance",
    "loan refinance savings": "refinance",
    "personal debt consolidation": "debt consolidation",
    "personal debt consolidator": "debt consolidation",
    "investment property": "rental property",
    "business loan": "commercial loan",
    "market value of bonds": "bond",
    "balloon mortgage": "balloon loan",
    "lump sum present value": "present value",
    "lump sum future value": "future value",
    "auto rebate vs low interest financing": "cash back or low interest",
    "low interest financing savings": "cash back or low interest",
    "mortgage with taxes and insurance": "mortgage",
    "mortgage with pmi": "mortgage",
    "fixed rate mortgage": "mortgage",
    "mortgage for purchase price and down payment piti": "mortgage for purchase price and down payment",
    "home equity line of credit": "heloc",
}

# Links on hub pages that are not calculators.
NON_CALCULATOR = re.compile(
    r"^(home|sign in|about|about us|contact|contact us|privacy|terms|sitemap|android|iphone/ipad|"
    r"other apps|all calculators|financial|fitness|math|other|more|help|feedback|support|news|"
    r"information|order|search)$", re.I)

# --- Workbook styling ---------------------------------------------------------
NAVY, BLUE, LIGHT, GREY = "1F3864", "2E75B6", "DDEBF7", "595959"
TABLE_STYLE = "TableStyleMedium2"
# Pastel fill per subcategory (cycled), so rows of one subcategory read as a band.
PASTELS = ["DDEBF7", "E2EFDA", "FCE4D6", "FFF2CC", "EDE2F6", "D9F2F2", "F8E1E7", "E7E6E6",
           "DEEAF1", "E4F3D8", "FBE5D6", "FFF5D1", "E9E1F5", "D5EFEF", "F6DCE3", "EDEDED",
           "D6E4F0", "E0EFD4", "FDE9D9", "FFF0C2", "E6DDF3", "DCEFF0", "F9E3EA", "F2F2F2"]


def soup(path):
    return BeautifulSoup(path.read_bytes(), "html.parser")


# Some dinkytown titles/link texts have a one-line description appended to the name.
DESCRIPTION_RE = re.compile(r"\s+(Helps|See|Estimate|Determine|Calculate|Find|Use|Compare|Enter|Learn)\b.*\.$")


def clean_name(text):
    text = re.sub(r"\s+", " ", text).strip()
    return DESCRIPTION_RE.sub("", text)


# US-only focus: Canadian calculators (dinkytown's Canadian section and "(Canadian)" versions)
# and tools for taxes the US doesn't have are excluded and listed on their own sheet.
NON_US_RE = re.compile(r"canad|\brrsp\b|\brrif\b|\btfsa\b|\bresp\b|\blira\b|\bvat\b|\bgst\b|\bhst\b", re.I)


def non_us_reason(r):
    if r["region"] == "Canada":
        return "Canadian"
    if NON_US_RE.search(r["name"]):
        return "Not applicable in the US"
    return ""


# Yearly editions of the same tool, e.g. "1040 Tax Calculator (Tax Year 2023)".
TAX_YEAR_RE = re.compile(r"\s*\((tax year \d{4}|prior tax year)\)", re.I)


def norm_key(name):
    """Normalised de-duplication key for a calculator name."""
    k = name.lower()
    k = k.replace("&", " and ").replace("’", "'")
    k = TAX_YEAR_RE.sub(" ", k)
    k = re.sub(r"(\d+)\s*\(([a-z])\)", r"\1\2", k)          # 401(k) -> 401k
    k = re.sub(r"\bpay off\b", "payoff", k)
    k = re.sub(r"\bcards\b", "card", k)
    k = re.sub(r"\bnestegg\b", "nest egg", k)
    k = re.sub(r"\((canadian|canada|us|u\.s\.)\)", " ", k)
    k = re.sub(r"\b(calculators?|calculator|calc|tool|estimator|planner|analysis|the|a|an|your|my)\b", " ", k)
    k = re.sub(r"\b(canadian|canada|us|u\.s\.)\b", " ", k)
    k = re.sub(r"[^a-z0-9 ]+", " ", k)
    k = re.sub(r"\s+", " ", k).strip()
    while ALIASES.get(k, k) != k:  # aliases may chain
        k = ALIASES[k]
    return k


def categorize(name):
    for rx, sub in OVERRIDES:
        if rx.search(name):
            return sub
    for sub, rx in FINANCE_SUBCATEGORIES:
        if rx.search(name):
            return sub
    return UNCATEGORIZED


# --- Per-site extraction ---------------------------------------------------
def extract_calculator_net():
    mirror = HERE / "calculator.net" / "mirror"
    hubs = {
        "financial-calculator.html": "Financial",
        "fitness-and-health-calculator.html": "Fitness & Health",
        "math-calculator.html": "Math",
        "other-calculator.html": "Other",
    }
    hub_files = set(hubs) | {"index.html", "about-us.html", "sitemap.html"}
    rows, seen = [], set()
    for hub, cat in hubs.items():
        s = soup(mirror / hub)
        for a in s.find_all("a", href=True):
            href = a["href"].split("#")[0]
            name = clean_name(a.get_text(" ", strip=True))
            if (not href.endswith(".html") or href in hub_files or "/" in href
                    or not name or NON_CALCULATOR.match(name) or href in seen
                    or not (mirror / href).exists()):
                continue
            seen.add(href)
            rows.append(dict(site="calculator.net", site_category=cat, name=name,
                             url=f"https://www.calculator.net/{href}", region="US / Global",
                             finance=(cat == "Financial")))
    return rows


def extract_dinkytown():
    mirror = HERE / "dinkytown.net" / "mirror"
    hubs = {
        "mortgage.html": "Mortgage", "loan.html": "Loan", "auto.html": "Auto",
        "debt.html": "Credit Cards & Debt", "investment.html": "Investment",
        "retirement.html": "Retirement", "savings.html": "Savings", "taxes.html": "Taxes",
        "insurance.html": "Insurance", "business.html": "Business", "personal.html": "Personal Finance",
        "payroll.html": "Payroll", "health-savings-account-hsa-calculators.html": "HSA",
        "ca.html": "Canadian",
        # Full listing last so hub categories win; catches anything not on a hub.
        "financialcalculators.html": "Full listing",
    }
    def hub_links(hub):
        path = mirror / hub
        if not path.exists():
            return set()
        return {a["href"].split("#")[0] for a in soup(path).find_all("a", href=True)}

    # French (FR.html) and Spanish (SP.html) pages are translations of English calculators.
    # Links in the shared site header appear on those pages too, so anything an English
    # category page also links to is not a translation.
    english = set().union(*(hub_links(h) for h in hubs if h != "financialcalculators.html"))
    translated = (hub_links("FR.html") | hub_links("SP.html")) - english
    rows, by_href = [], {}
    for hub, cat in hubs.items():
        path = mirror / hub
        if not path.exists():
            continue
        for a in soup(path).find_all("a", href=True):
            href = a["href"].split("#")[0]
            if not href.startswith("java/") or not href.endswith(".html") or href in translated:
                continue
            if href in by_href:
                # Record every hub the calculator appears on.
                if cat not in by_href[href]["site_category"].split(" / ") and cat != "Full listing":
                    by_href[href]["site_category"] += f" / {cat}"
                continue
            target = mirror / href
            if target.exists():
                # <title> is the bare name; <h1> sometimes has a description appended.
                t = soup(target)
                h1 = t.find("h1")
                name = clean_name((t.title.string if t.title and t.title.string else "")
                                  or (h1.get_text(" ", strip=True) if h1 else ""))
            else:
                name = clean_name(a.get_text(" ", strip=True))
            if not name:
                continue
            canadian = "canadian" in href.lower() or "(canadian)" in name.lower() or cat == "Canadian"
            row = dict(site="dinkytown.net", site_category=cat, name=name,
                       url=f"https://www.dinkytown.net/{href}",
                       region="Canada" if canadian else "US", finance=True)
            by_href[href] = row
            rows.append(row)
    return rows


def extract_fncalculator():
    mirror = HERE / "fncalculator.com" / "mirror"
    rows, seen = [], set()
    for a in soup(mirror / "index.html").find_all("a", href=True):
        href = a["href"]
        name = clean_name(a.get_text(" ", strip=True))
        if not href.startswith(("financialcalculator__type-", "currencyConverter")):
            continue
        if not name or NON_CALCULATOR.match(name) or href in seen:
            continue
        seen.add(href)
        m = re.match(r"financialcalculator__type-(.+)\.html", href)
        live = f"https://www.fncalculator.com/financialcalculator?type={m.group(1)}" if m \
            else f"https://www.fncalculator.com/{href.removesuffix('.html')}"
        rows.append(dict(site="fncalculator.com", site_category="", name=name, url=live,
                         region="US" if name.lower().startswith("us ") else "US / Global", finance=True))
    return rows


# --- Workbook --------------------------------------------------------------
def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def add_table(ws, name, headers, data, widths, first_row=1, style=TABLE_STYLE):
    """Write headers + data starting at first_row and format them as a banded Excel table."""
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=first_row, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r, row in enumerate(data, first_row + 1):
        for c, value in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    last_row = first_row + max(len(data), 1)
    ref = f"A{first_row}:{get_column_letter(len(headers))}{last_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(table)
    ws.row_dimensions[first_row].height = 32
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return first_row + 1, last_row


def link(cell, url, text="Open ↗"):
    if url:
        cell.value, cell.hyperlink = text, url
        cell.font = Font(color="0563C1", underline="single")


def sheet_title(ws, title, subtitle, width_cols):
    last = get_column_letter(width_cols)
    ws.merge_cells(f"A1:{last}1")
    ws.merge_cells(f"A2:{last}2")
    ws["A1"] = title
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = fill(NAVY)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws["A2"] = subtitle
    ws["A2"].font = Font(size=10, italic=True, color=GREY)
    ws["A2"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 20
    ws.sheet_view.showGridLines = False


def build_workbook(rows, unique, excluded):
    wb = Workbook()
    subcats = [n for n, _ in FINANCE_SUBCATEGORIES] + [UNCATEGORIZED]
    used = [n for n in subcats if any(u["subcategory"] == n for u in unique)]
    colour = {n: PASTELS[i % len(PASTELS)] for i, n in enumerate(used)}
    today = date.today().strftime("%-d %B %Y")
    source = f"US calculators only  ·  Sources: {', '.join(SITES)}  ·  Generated {today}"

    # --- Overview ---------------------------------------------------------
    ws = wb.active
    ws.title = "Overview"
    sheet_title(ws, "Competitor Finance Calculator Inventory (US)", source, 9)

    kpis = [
        ("Unique finance calculators", len(unique)),
        ("Subcategories", len(used)),
        ("On all 3 sites", sum(1 for u in unique if u["site_count"] == 3)),
        ("On only 1 site", sum(1 for u in unique if u["site_count"] == 1)),
    ]
    thin = Side(style="thin", color="BFBFBF")
    for i, (label, value) in enumerate(kpis):
        col = 1 + i * 2
        ws.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 1)
        ws.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col + 1)
        v = ws.cell(row=4, column=col, value=value)
        v.font = Font(size=24, bold=True, color=BLUE)
        lab = ws.cell(row=5, column=col, value=label)
        lab.font = Font(size=10, color=GREY)
        for r in (4, 5):
            for c in (col, col + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = fill(LIGHT)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = Border(top=thin if r == 4 else None, bottom=thin if r == 5 else None,
                                     left=thin if c == col else None, right=thin if c == col + 1 else None)
    ws.row_dimensions[4].height = 40
    ws.row_dimensions[5].height = 22

    # Per-site table
    ws["A7"] = "Coverage by site"
    ws["A7"].font = Font(size=12, bold=True, color=NAVY)
    site_data = []
    for s in SITES:
        src = [r for r in rows if r["site"] == s]
        fin = [r for r in src if r["finance"]]
        on_site = [u for u in unique if u["per_site"][s]]
        exclusive = [u for u in on_site if u["site_count"] == 1]
        site_data.append([s, len(src), len(fin), len(on_site), len(exclusive)])
    add_table(ws, "SiteCoverage",
              ["Site", "Calculators listed", "Finance calculators", "Unique finance tools covered",
               "Only on this site"], site_data, [20, 14, 14, 14, 14, 14, 14, 14, 14], first_row=8)

    # Subcategory summary
    start = 8 + len(site_data) + 3
    ws.cell(row=start - 1, column=1, value="Finance calculators by subcategory").font = \
        Font(size=12, bold=True, color=NAVY)
    sum_data = []
    for n in used:
        us = [u for u in unique if u["subcategory"] == n]
        sum_data.append([n, len(us)] + [sum(1 for u in us if u["per_site"][s]) for s in SITES] +
                        [sum(1 for u in us if u["site_count"] == 3), sum(1 for u in us if u["site_count"] == 1)])
    sum_data.sort(key=lambda r: -r[1])
    first, last = add_table(ws, "SubcategorySummary",
                            ["Subcategory", "Unique calculators"] + SITES + ["On all 3 sites", "On only 1 site"],
                            sum_data, [44, 14, 14, 14, 16, 14, 14], first_row=start)
    for r in range(first, last + 1):
        ws.cell(row=r, column=1).fill = fill(colour[ws.cell(row=r, column=1).value])
        for c in range(2, 8):
            ws.cell(row=r, column=c).alignment = Alignment(horizontal="center", vertical="top")
    ws.conditional_formatting.add(f"B{first}:B{last}", DataBarRule(start_type="num", start_value=0,
                                  end_type="max", color="5B9BD5", showValue=True))
    ws.conditional_formatting.add(f"C{first}:E{last}", ColorScaleRule(start_type="min", start_color="FFFFFF",
                                  end_type="max", end_color="9BC2E6"))
    total_row = last + 1
    ws.cell(row=total_row, column=1, value="Total")
    ws.cell(row=total_row, column=2, value=len(unique))
    for i, s in enumerate(SITES, 3):
        ws.cell(row=total_row, column=i, value=sum(1 for u in unique if u["per_site"][s]))
    ws.cell(row=total_row, column=6, value=kpis[2][1])
    ws.cell(row=total_row, column=7, value=kpis[3][1])
    for c in range(1, 8):
        cell = ws.cell(row=total_row, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill(NAVY)
        cell.alignment = Alignment(horizontal="left" if c == 1 else "center")

    # Bar chart of calculators per subcategory
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "Unique calculators per subcategory"
    chart.y_axis.title = None
    chart.x_axis.title = None
    chart.legend = None
    chart.add_data(Reference(ws, min_col=2, min_row=first - 1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
    chart.x_axis.scaling.orientation = "maxMin"
    chart.series[0].graphicalProperties.solidFill = "2E75B6"
    chart.height, chart.width = 13, 18
    ws.add_chart(chart, f"I{start}")

    # --- Calculators (unique, finance) -----------------------------------------
    ws = wb.create_sheet("Calculators")
    headers = ["Subcategory", "Calculator", "Sites (of 3)"] + SITES + \
              ["Name variants on other sites"] + [f"{s} link" for s in SITES] + ["Other URLs"]
    sheet_title(ws, "Unique finance calculators", f"{len(unique)} calculators merged across sites  ·  "
                "✓ = the site has it  ·  links open the competitor page  ·  " + source, len(headers))
    data = []
    for u in unique:
        extra = [m["url"] for s in SITES for m in u["per_site"][s][1:]]
        data.append([u["subcategory"], u["name"], u["site_count"]] +
                     ["✓" if u["per_site"][s] else "" for s in SITES] +
                     ["; ".join(u["variants"])] + [""] * len(SITES) + ["\n".join(extra)])
    first, last = add_table(ws, "Calculators", headers, data,
                            [36, 40, 10, 13, 13, 15, 44, 13, 13, 15, 50], first_row=4)
    for i, u in enumerate(unique):
        r = first + i
        ws.cell(row=r, column=1).fill = fill(colour[u["subcategory"]])
        ws.cell(row=r, column=2).font = Font(bold=True)
        for c in range(3, 7):
            ws.cell(row=r, column=c).alignment = Alignment(horizontal="center", vertical="top")
        for j, s in enumerate(SITES):
            if u["per_site"][s]:
                link(ws.cell(row=r, column=8 + j), u["per_site"][s][0]["url"])
                ws.cell(row=r, column=8 + j).alignment = Alignment(horizontal="center", vertical="top")
    ws.conditional_formatting.add(f"D{first}:F{last}", CellIsRule(operator="equal", formula=['"✓"'],
                                  fill=fill("C6EFCE"), font=Font(color="006100", bold=True)))
    ws.conditional_formatting.add(f"C{first}:C{last}", ColorScaleRule(
        start_type="num", start_value=1, start_color="FCE4D6",
        mid_type="num", mid_value=2, mid_color="FFF2CC",
        end_type="num", end_value=3, end_color="C6EFCE"))
    ws.freeze_panes = "C5"

    # --- All source rows ---------------------------------------------------------
    ws = wb.create_sheet("All source rows")
    headers = ["Site", "Site's own category", "Calculator", "Type", "Subcategory", "Match key", "URL"]
    sheet_title(ws, "Every calculator as listed on each site",
                "One row per calculator per site, before merging  ·  Match key is what merges rows  ·  " + source,
                len(headers))
    data = [[r["site"], r["site_category"], r["name"], "Finance" if r["finance"] else "Non-finance",
             r["subcategory"], r["key"], ""] for r in rows]
    first, _ = add_table(ws, "SourceRows", headers, data, [18, 26, 42, 12, 36, 30, 13],
                         first_row=4, style="TableStyleMedium9")
    for i, r in enumerate(rows):
        link(ws.cell(row=first + i, column=7), r["url"])
        if r["subcategory"]:
            ws.cell(row=first + i, column=5).fill = fill(colour[r["subcategory"]])
    ws.freeze_panes = "D5"

    # --- Non-finance (reference) -------------------------------------------------
    ws = wb.create_sheet("Non-finance (reference)")
    nonfin = [r for r in rows if not r["finance"]]
    headers = ["Site's category", "Calculator", "URL"]
    sheet_title(ws, "Non-finance calculators (calculator.net)",
                "Kept for reference; outside the finance subcategories  ·  " + source, len(headers))
    first, _ = add_table(ws, "NonFinance", headers, [[r["site_category"], r["name"], ""] for r in nonfin],
                         [22, 44, 13], first_row=4, style="TableStyleMedium7")
    for i, r in enumerate(nonfin):
        link(ws.cell(row=first + i, column=3), r["url"])
    ws.freeze_panes = "A5"

    # --- Excluded (non-US) ---------------------------------------------------------
    ws = wb.create_sheet("Excluded (non-US)")
    headers = ["Reason", "Site", "Site's own category", "Calculator", "URL"]
    sheet_title(ws, "Excluded: non-US calculators",
                f"{len(excluded)} calculators left out of the US inventory  ·  " + source, len(headers))
    first, _ = add_table(ws, "Excluded", headers,
                         [[r["reason"], r["site"], r["site_category"], r["name"], ""] for r in excluded],
                         [22, 18, 26, 50, 13], first_row=4, style="TableStyleMedium3")
    for i, r in enumerate(excluded):
        link(ws.cell(row=first + i, column=5), r["url"])
    ws.freeze_panes = "A5"

    for sheet in wb.worksheets:
        sheet.sheet_properties.tabColor = {"Overview": NAVY, "Calculators": BLUE}.get(sheet.title, "A6A6A6")
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(OUT_FILE)


def main():
    extractors = {
        "calculator.net": extract_calculator_net,
        "dinkytown.net": extract_dinkytown,
        "fncalculator.com": extract_fncalculator,
    }
    rows = []
    for site, fn in extractors.items():
        site_rows = fn()
        print(f"{site}: {len(site_rows)} calculators")
        rows.extend(site_rows)

    excluded = [dict(r, reason=non_us_reason(r)) for r in rows if non_us_reason(r)]
    rows = [r for r in rows if not non_us_reason(r)]
    print(f"Excluded as non-US: {len(excluded)}")

    for r in rows:
        r["key"] = norm_key(r["name"])
        r["subcategory"] = categorize(r["name"]) if r["finance"] else ""

    # Merge the same calculator across sites.
    groups = defaultdict(list)
    for r in rows:
        if r["finance"]:
            groups[r["key"]].append(r)

    unique = []
    for key, members in groups.items():
        # Display name: the plainest naming (calculator.net, then fncalculator, then dinkytown).
        site_rank = {"calculator.net": 0, "fncalculator.com": 1, "dinkytown.net": 2}
        display = sorted(members, key=lambda r: (site_rank[r["site"]], len(r["name"])))[0]
        per_site = {s: [m for m in members if m["site"] == s] for s in SITES}
        unique.append(dict(
            subcategory=display["subcategory"],
            name=TAX_YEAR_RE.sub("", display["name"]).strip(),
            key=key,
            site_count=sum(1 for s in SITES if per_site[s]),
            per_site=per_site,
            variants=sorted({m["name"] for m in members} - {display["name"]}),
        ))
    unique.sort(key=lambda u: (u["subcategory"], -u["site_count"], u["name"].lower()))

    build_workbook(rows, unique, excluded)
    uncategorized = [u["name"] for u in unique if u["subcategory"] == UNCATEGORIZED]
    print(f"Unique finance calculators: {len(unique)}; subcategories used: "
          f"{len({u['subcategory'] for u in unique})}; uncategorized: {len(uncategorized)}")
    if uncategorized:
        print("Uncategorized:", "; ".join(uncategorized))
    print(f"Saved {OUT_FILE}")


if __name__ == "__main__":
    main()

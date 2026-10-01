#!/usr/bin/env python3
"""
One-off seeding: create src/content/tools/<slug>.json stubs for every calculator in
utilities/competitor_research/calculator_inventory.xlsx, plus the 8 calculators from the old
site in /archive, and stubs for the info pages in src/content/pages/.

Never overwrites an existing file, so it is safe to re-run after the inventory grows: only new
calculators get stubs. Once a tool's JSON exists, that file is the source of truth — edit it,
not this script.

Run with the competitor_research venv (needs openpyxl), from the repo root:
    utilities/competitor_research/.venv/bin/python utilities/scaffold/seed_tool_stubs.py
"""
import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
INVENTORY = ROOT / "utilities" / "competitor_research" / "calculator_inventory.xlsx"
CATEGORIES = ROOT / "src" / "config" / "categories.json"
TOOLS_DIR = ROOT / "src" / "content" / "tools"
PAGES_DIR = ROOT / "src" / "content" / "pages"
SITES = ["calculator.net", "dinkytown.net", "fncalculator.com"]

# Names that already say what kind of tool they are; everything else gets "-calculator" appended.
TOOL_NOUN_RE = re.compile(r"calculator|planner|estimator|converter|analyzer|optimizer|worksheet|"
                          r"questionnaire|balancer|comparison|analysis")

# The old MarketCapInsights calculators keep their existing, already-indexed URLs as slugs.
ARCHIVE_TOOLS = [
    ("market-cap-growth-calculator", "Market Cap Growth Calculator", "stock-calculators",
     "index.html", "/"),
    ("position-sizing", "Position Sizing Calculator", "stock-calculators",
     "position-sizing.html", "/position-sizing.html"),
    ("tax-harvesting", "Tax-Efficient Harvesting Calculator", "tax-calculators",
     "tax-harvesting.html", "/tax-harvesting.html"),
    ("sector-balance", "Portfolio Sector Balance Optimizer", "investment-calculators",
     "sector-balance.html", "/sector-balance.html"),
    ("stock-split-impact", "Stock Split Impact Calculator", "stock-calculators",
     "stock-split-impact.html", "/stock-split-impact.html"),
    ("dividend-reinvestment", "Dividend Reinvestment Breakeven Calculator", "stock-calculators",
     "dividend-reinvestment.html", "/dividend-reinvestment.html"),
    ("stock-valuation-confidence", "Stock Valuation Confidence Interval Calculator", "stock-calculators",
     "stock-valuation-confidence.html", "/stock-valuation-confidence.html"),
    ("inflation-adjusted-calculator", "Inflation-Adjusted Return Calculator", "investment-calculators",
     "inflation-adjusted-calculator.html", "/inflation-adjusted-calculator.html"),
]

INFO_PAGES = [
    ("about", "About MarketCapInsights", "about.html"),
    ("contact", "Contact Us", "contact.html"),
    ("privacy-policy", "Privacy Policy", "privacy-policy.html"),
    ("terms", "Terms and Conditions", "terms.html"),
]


def slugify(name):
    s = name.lower().replace("&", " and ").replace("'", "").replace("’", "")
    s = re.sub(r"(\d+)\s*\(([a-z])\)", r"\1\2", s)          # 401(k) -> 401k
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    s = re.sub(r"-+", "-", s)
    if not TOOL_NOUN_RE.search(s):
        s += "-calculator"
    return s


def stub(slug, name, subcategory, priority, order, research=None, legacy=None):
    data = {
        "slug": slug,
        "status": "planned",
        "name": name,
        "nav_name": re.sub(r"\s+Calculator$", "", name) or name,
        "subcategory": subcategory,
        "priority": priority,
        "order": order,
        "meta_title": "",
        "meta_description": "",
        "h1": name,
        "subtitle": "",
        "card": {"layout": "raw", "fields_html": "", "extra_scripts": []},
        "script": "",
        "content_html": "",
        "faq": [],
        "related": [],
    }
    if legacy:
        data["legacy"] = legacy
    if research:
        data["research"] = research
    return data


def write_new(path, data):
    if path.exists():
        return False
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return True


def main():
    TOOLS_DIR.mkdir(parents=True, exist_ok=True)
    PAGES_DIR.mkdir(parents=True, exist_ok=True)
    cats = json.loads(CATEGORIES.read_text())
    sub_by_name = {s["name"]: s["slug"] for g in cats["nav_groups"] for s in g["subcategories"]}

    ws = load_workbook(INVENTORY)["Calculators"]
    header = [c.value for c in ws[4]]
    col = {h: i for i, h in enumerate(header)}
    created, skipped, seen, per_sub = 0, 0, set(), {}

    for row in ws.iter_rows(min_row=5):
        name = row[col["Calculator"]].value
        if not name:
            continue
        sub_name = row[col["Subcategory"]].value
        subcategory = sub_by_name[sub_name]
        slug = slugify(name)
        if slug in seen:
            raise SystemExit(f"Duplicate slug {slug!r} from {name!r} — rename one in the inventory aliases")
        seen.add(slug)
        per_sub[subcategory] = per_sub.get(subcategory, 0) + 1
        competitors = []
        for site in SITES:
            cell = row[col[f"{site} link"]]
            if cell.hyperlink and cell.hyperlink.target:
                competitors.append({"site": site, "url": cell.hyperlink.target})
        variants = row[col["Name variants on other sites"]].value
        research = {
            "competitor_count": row[col["Sites (of 3)"]].value,
            "competitors": competitors,
            "name_variants": [v.strip() for v in variants.split(";")] if variants else [],
        }
        data = stub(slug, name, subcategory, priority=row[col["Sites (of 3)"]].value,
                    order=per_sub[subcategory] * 10, research=research)
        if write_new(TOOLS_DIR / f"{slug}.json", data):
            created += 1
        else:
            skipped += 1

    for slug, name, subcategory, source, url in ARCHIVE_TOOLS:
        if slug in seen:
            raise SystemExit(f"Archive slug {slug!r} collides with an inventory slug")
        seen.add(slug)
        legacy = {"source": f"archive/{source}", "url": url,
                  "note": "Existing MarketCapInsights calculator: port its logic into card/script "
                          "and rewrite its copy into this file."}
        # Archive tools already have traffic, so they sort first in their hub.
        data = stub(slug, name, subcategory, priority=3, order=1, legacy=legacy)
        if write_new(TOOLS_DIR / f"{slug}.json", data):
            created += 1
        else:
            skipped += 1

    pages_created = 0
    for slug, h1, source in INFO_PAGES:
        page = {
            "slug": slug,
            "status": "planned",
            "meta_title": "",
            "meta_description": "",
            "h1": h1,
            "subtitle": "",
            "content_html": "",
            "legacy": {"source": f"archive/{source}", "url": f"/{source}",
                       "note": "Rewrite for the new site; review AdSense/privacy wording."},
        }
        pages_created += write_new(PAGES_DIR / f"{slug}.json", page)

    print(f"Tool stubs: {created} created, {skipped} already existed ({len(seen)} tools total)")
    print(f"Info page stubs: {pages_created} created")


if __name__ == "__main__":
    main()
